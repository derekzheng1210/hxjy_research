# -*- coding: utf-8 -*-
"""DM 信用债盘面日报：新旧附件格式解析与游标分页增量更新。"""
from __future__ import annotations

import json
import os
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

# 先于 broker_market.paths 解析：默认沙箱，避免组合跑测试时把模块常量
# 绑定到生产数据目录。
os.environ.setdefault("PORTAL_DATA_ROOT", str(Path(__file__).resolve().parent / ".test_runtime" / "dm_report"))

from broker_market import dm_report


def _row(ws, row, values, start_col=1):
    for offset, value in enumerate(values):
        ws.cell(row, start_col + offset).value = value


def _make_panorama_workbook(path: Path):
    """新格式：全景图 sheet，经纪商 5 行块（收益率/日变动/笔数/TKN/GVN）+ 平台块。"""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "全景图"
    _row(ws, 2, ["10月28日信用债成交收益全景图"])
    _row(ws, 3, ["来源", "券种细分", None, "剩余期限"])
    _row(ws, 4, [None, None, None, "<1Y", "1-3Y", "3-5Y", "5-7Y", "7-10Y", ">10Y"])
    _row(ws, 5, ["经纪商", "银行二永债", "收益率", 1.4, 1.5, 1.6, None, None, None])
    _row(ws, 6, [None, None, "日变动(BP)", 1.0, -1.0, 0.8, None, None, None])
    _row(ws, 7, [None, None, "成交笔数", 124, 137, 493, 26, 79, None])
    _row(ws, 8, [None, None, "TKN(%)", 0.6, 0.55, 0.58, 0.4, 0.6, None])
    _row(ws, 9, [None, None, "GVN(%)", 0.4, 0.45, 0.42, 0.6, 0.4, None])
    _row(ws, 10, [None, "产业债", "收益率", 1.5, 1.6, None, None, None, None])
    _row(ws, 11, [None, None, "日变动(BP)", 2.0, 1.0, None, None, None, None])
    _row(ws, 12, [None, None, "成交笔数", 171, 265, None, None, None, None])
    _row(ws, 13, [None, None, "TKN(%)", 0.7, 0.65, None, None, None, None])
    _row(ws, 14, [None, None, "GVN(%)", 0.3, 0.35, None, None, None, None])
    _row(ws, 15, [None, "城投债", "收益率", 1.6, 1.7, None, None, None, None])
    _row(ws, 16, [None, None, "日变动(BP)", 2.0, 1.0, None, None, None, None])
    _row(ws, 17, [None, None, "成交笔数", 101, 253, None, None, None, None])
    _row(ws, 18, [None, None, "TKN(%)", 0.6, 0.6, None, None, None, None])
    _row(ws, 19, [None, None, "GVN(%)", 0.4, 0.4, None, None, None, None])
    # 平台块：来源不是经纪商，必须被排除
    _row(ws, 20, ["平台", "银行二永债", "收益率", 1.4, 1.5, None, None, None, None])
    _row(ws, 21, [None, None, "成交笔数", 999, 999, None, None, None, None])
    _row(ws, 22, [None, None, "TKN(%)", 0.9, 0.9, None, None, None, None])
    wb.save(path)
    wb.close()


def _make_emotion_workbook(path: Path):
    """旧格式：独立成交情绪 sheet，券种块标签在 A 列、指标在 B 列。"""
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "成交情绪 "
    _row(ws, 1, ["本日经纪商来源成交类型分布"])
    _row(ws, 2, ["券种细分", None, "剩余期限"])
    _row(ws, 3, [None, None, "<1Y", "1-3Y", "3-5Y", "5-7Y", "7-10Y", ">10Y"])
    _row(ws, 4, ["全部", "成交笔数", 747, 1252, 1101, 32, 231, 17])
    _row(ws, 5, [None, "TKN(%)", 0.75, 0.69, 0.64, 0.53, 0.63, 0.59])
    _row(ws, 6, [None, "GVN(%)", 0.23, 0.29, 0.35, 0.44, 0.36, 0.41])
    _row(ws, 7, ["银行二永债", "成交笔数", 102, 297, 782, 26, 146, None])
    _row(ws, 8, [None, "TKN(%)", 0.75, 0.63, 0.60, 0.5, 0.60, None])
    _row(ws, 9, [None, "GVN(%)", 0.25, 0.37, 0.38, 0.46, 0.40, None])
    _row(ws, 10, ["券商债", "成交笔数", 52, 71, 7, None, None, None])
    _row(ws, 11, [None, "TKN(%)", 0.77, 0.73, 0.86, None, None, None])
    _row(ws, 12, [None, "GVN(%)", 0.23, 0.27, 0.14, None, None, None])
    _row(ws, 13, ["产业债", "成交笔数", 235, 348, 88, None, 70, 16])
    _row(ws, 14, [None, "TKN(%)", 0.77, 0.73, 0.68, None, 0.66, 0.56])
    _row(ws, 15, [None, "GVN(%)", 0.23, 0.27, 0.32, None, 0.30, 0.44])
    wb.save(path)
    wb.close()


class ParseReportTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(dir=Path.cwd() / '.test_runtime')
        self.root = Path(self.tmp.name)

    def tearDown(self):
        self.tmp.cleanup()

    def test_parse_panorama_new_format(self):
        path = self.root / "report_new.xlsx"
        _make_panorama_workbook(path)
        record = dm_report.parse_report_file(path, "2026-10-28")
        self.assertEqual(record["day"], "2026-10-28")
        # 859 = 124+137+493+26+79；436 = 171+265；354 = 101+253
        self.assertEqual(record["num"], 859 + 436 + 354)
        er = record["sectors"]["银行二永债"]
        self.assertEqual(er["num"], 859)
        self.assertEqual(er["tkn"], round(124 * 0.6) + round(137 * 0.55) + round(493 * 0.58) + round(26 * 0.4) + round(79 * 0.6))
        self.assertEqual(er["gvn"], round(124 * 0.4) + round(137 * 0.45) + round(493 * 0.42) + round(26 * 0.6) + round(79 * 0.4))
        # 平台块 999 笔必须被排除；城投块 101+253=354
        self.assertEqual(record["num"], 859 + 436 + 354)
        self.assertEqual(record["sectors"]["产业债"]["maturity"]["1-3Y"]["tkn_pct"], 65.0)

    def test_parse_emotion_old_format(self):
        path = self.root / "report_old.xlsx"
        _make_emotion_workbook(path)
        record = dm_report.parse_report_file(path, "2025-11-05")
        self.assertEqual(record["day"], "2025-11-05")
        self.assertEqual(record["sectors"]["银行二永债"]["num"], 102 + 297 + 782 + 26 + 146)
        self.assertEqual(record["sectors"]["产业债"]["num"], 235 + 348 + 88 + 70 + 16)
        self.assertNotIn("全部", record["sectors"])
        self.assertEqual(record["num"],
                         record["sectors"]["银行二永债"]["num"] + record["sectors"]["产业债"]["num"]
                         + record["sectors"]["券商债"]["num"])

    def test_unrecognized_file_returns_none(self):
        import openpyxl

        path = self.root / "blank.xlsx"
        wb = openpyxl.Workbook(); wb.save(path); wb.close()
        self.assertIsNone(dm_report.parse_report_file(path, "2026-10-28"))

    def test_platform_pivot_sheet_not_confused_with_broker_emotion(self):
        """2026-01-20 实际文件：含平台来源透视表（名称也带“成交情绪”），
        必须选经纪商口径的「成交情绪」sheet 解析。"""
        import openpyxl

        path = self.root / "report_pivot.xlsx"
        _make_emotion_workbook(path)
        wb = openpyxl.load_workbook(path)
        # 在成交情绪 sheet 前插入一个平台来源透视表 sheet
        pivot = wb.create_sheet("（自动表）自动计算平台来源成交情绪 ", 0)
        _row(pivot, 4, [None, "<1Y", None, None, ">10Y", None, None, "1-3Y", None, None,
                        "3-5Y", None, None, "5-7Y", None, None, "7-10Y"])
        _row(pivot, 5, ["行标签", "求和项:总成交笔数", "求和项:TKN", "求和项:GVN"])
        _row(pivot, 11, ["商金债", 110, 0, 0])
        wb.save(path)
        wb.close()
        record = dm_report.parse_report_file(path, "2026-01-20")
        self.assertIsNotNone(record)
        # 透视表数值（如 110）不得进入结果；以经纪商成交情绪 sheet 为准
        self.assertEqual(record["num"], 1353 + 130 + 757)
        self.assertIn("银行二永债", record["sectors"])
        self.assertGreater(record["tkn"], 0)


class _Response:
    def __init__(self, payload):
        self._payload = payload
        self.status_code = 200

    def raise_for_status(self):
        pass

    def json(self):
        return self._payload


class _Session:
    def __init__(self, pages, files):
        self._pages = list(pages)
        self._files = files
        self.posts = 0

    def post(self, url, data=None, headers=None, timeout=None):
        self.posts += 1
        return _Response(self._pages.pop(0) if self._pages else {"code": 0, "data": {"list": []}})

    def get(self, url, timeout=None):
        return _FileResponse(self._files[url])


class _FileResponse:
    def __init__(self, content):
        self.content = content
        self.status_code = 200

    def raise_for_status(self):
        pass


class _Client:
    def __init__(self, pages, files):
        self.session = _Session(pages, files)
        self.timeout = 10

    def signed_url(self, path, **extra):
        return path if not extra else f"{path}?{extra}"


class UpdateHistoryTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(dir=Path.cwd() / '.test_runtime')
        self.root = Path(self.tmp.name)
        self.patches = [
            patch.object(dm_report, "HISTORY_PATH", self.root / "history.json"),
            patch.object(dm_report, "REPORTS_DIR", self.root / "dm_reports"),
        ]
        for p in self.patches:
            p.start()

    def tearDown(self):
        for p in reversed(self.patches):
            p.stop()
        self.tmp.cleanup()

    def _workbook_bytes(self):
        path = self.root / "seed.xlsx"
        _make_panorama_workbook(path)
        return path.read_bytes()

    def test_cursor_pagination_and_incremental_update(self):
        article = {"title": "信用债盘面日报1028 | 测试", "issueTime": 1793180384000,
                   "sentimentId": "1",
                   "annexList": [{"sentimentFileUrl": "https://example.com/a.xlsx",
                                  "sentimentFileName": "DM-信用债盘面日报1028.xlsx"}]}
        # 第 1 页带游标翻页信息，第 2 页收尾
        pages = [
            {"code": 0, "data": {"list": [article], "hasNextPage": True,
                                 "endPageId": "1786440194064692143"}},
            {"code": 0, "data": {"list": [], "hasNextPage": False, "endPageId": ""}},
        ]
        client = _Client(pages, {"https://example.com/a.xlsx": self._workbook_bytes()})
        result = dm_report.update_history(client, start="2026-10-28", end="2026-10-28")
        self.assertEqual(result["fetched"], 1)
        history = dm_report.load_history()
        self.assertIn("2026-10-28", history["days"])
        self.assertEqual(history["days"]["2026-10-28"]["num"], 859 + 436 + 354)
        # 游标分页必须翻到第 2 页（共 2 次 POST）
        self.assertEqual(client.session.posts, 2)
        # 已有日期默认跳过：再次全量更新不重复下载
        client2 = _Client([{"code": 0, "data": {"list": [article], "hasNextPage": False}}], {})
        result2 = dm_report.update_history(client2, start="2026-10-28", end="2026-10-28")
        self.assertEqual(result2["fetched"], 0)

    def test_history_needs_update_gate(self):
        from datetime import date, timedelta
        yesterday = (date.today() - timedelta(days=1)).isoformat()
        dm_report.save_history({"days": {yesterday: {"num": 1}}})
        self.assertFalse(dm_report.history_needs_update())
        dm_report.save_history({"days": {"2026-01-01": {"num": 1}}})
        self.assertTrue(dm_report.history_needs_update())
        dm_report.save_history({"days": {}})
        self.assertTrue(dm_report.history_needs_update())


if __name__ == "__main__":
    unittest.main()
