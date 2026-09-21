# -*- coding: utf-8 -*-
"""机构行为净买入 · 最细颗粒度全量 Excel 导出

数据源：机构行为监测上游 bondflow API（与 /institution-flow 页面同源）。
口径：日度净买入（亿元），正值=净买入，负值=净卖出。
颗粒度：交易日 × 券种 × 期限 × 机构（全量历史，2019-01-01 起）。

用法：
    .venv/Scripts/python.exe scripts/export_institution_flow_excel.py
输出：
    D:/信用债研究/机构行为净买入_分券种分期限分机构_全量_<生成日>.xlsx
    （每个券种一个 sheet：行=交易日，列=期限×机构 + 当日合计）
"""
import json
import sys
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from pathlib import Path

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

UPSTREAM = "http://43.137.12.140:8000/bondflow/api"
DATA_START = "2019-01-01"
WORKERS = 6
RETRIES = 3
OUT_DIR = Path(r"D:\信用债研究")

# 展示顺序：期限短→长；机构按 配置盘→交易盘→其他；券种按 利率债→信用债→存单及其他
TENOR_ORDER = ["1年及1年以下", "1-3年", "3-5年", "5-7年", "7-10年",
               "10-15年", "15-20年", "20-30年", "30年以上"]
INST_ORDER = ["大型银行", "中小型银行", "保险公司", "理财子公司及理财类产品",
              "基金公司及产品", "证券公司", "货币市场基金", "其他"]
BOND_ORDER = ["国债", "政金债", "地方政府债", "中期票据", "企业债", "短期和超短期融资券",
              "同业存单", "资产支持证券", "其他"]

session = requests.Session()
session.headers.update({"Referer": "http://43.137.12.140:8000/jgxw/"})


def fetch_options():
    r = session.get(f"{UPSTREAM}/options/", timeout=60)
    r.raise_for_status()
    return r.json()


def fetch_combo(bond, tenor, end_date):
    """单券种×期限，按机构分解，返回 {inst: {date: value}}"""
    params = [
        ("institutions", inst) for inst in INST_ORDER
    ] + [
        ("bond_types", bond), ("tenors", tenor),
        ("start_date", DATA_START), ("end_date", end_date),
        ("granularity", "day"), ("dimension", "institution"),
    ]
    last_err = None
    for attempt in range(RETRIES):
        try:
            r = session.get(f"{UPSTREAM}/dimension/", params=params, timeout=120)
            r.raise_for_status()
            out = {}
            for item in r.json().get("data") or []:
                name = item.get("name")
                series = {}
                for p in item.get("series") or []:
                    v = p.get("net_buy", p.get("value"))
                    if v is not None:
                        series[p["date"]] = round(float(v), 2)
                if name and series:
                    out[name] = series
            return bond, tenor, out
        except Exception as e:  # noqa: BLE001
            last_err = e
            time.sleep(1.5 * (attempt + 1))
    raise RuntimeError(f"{bond} × {tenor} 拉取失败：{last_err}")


def build_sheet(wb, bond, data, all_dates):
    """一个券种一个 sheet：行=交易日，列=期限×机构 + 当日合计"""
    ws = wb.create_sheet(bond)

    head_fill = PatternFill("solid", fgColor="1F4E79")
    head_font = Font(bold=True, color="FFFFFF", size=10)
    sub_fill = PatternFill("solid", fgColor="DDEBF7")
    sub_font = Font(bold=True, size=10)
    thin = Side(style="thin", color="B0B0B0")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    center = Alignment(horizontal="center", vertical="center")

    n_tenor, n_inst = len(TENOR_ORDER), len(INST_ORDER)
    last_col = 1 + n_tenor * n_inst + 1  # 日期 + 数据列 + 合计

    # 表头两行：期限（合并）/ 机构
    ws.cell(row=1, column=1, value="日期")
    ws.cell(row=2, column=1, value="日期")
    for ti, tenor in enumerate(TENOR_ORDER):
        c0 = 2 + ti * n_inst
        cell = ws.cell(row=1, column=c0, value=tenor)
        if n_inst > 1:
            ws.merge_cells(start_row=1, start_column=c0, end_row=1, end_column=c0 + n_inst - 1)
        for ci in range(n_inst):
            ws.cell(row=2, column=c0 + ci, value=INST_ORDER[ci])
    total_c = last_col
    ws.cell(row=1, column=total_c, value="当日合计")
    ws.merge_cells(start_row=1, start_column=total_c, end_row=2, end_column=total_c)
    for c in range(1, last_col + 1):
        is_edge = (c == 1 or c == total_c)
        fill = head_fill if is_edge else sub_fill
        font = head_font if is_edge else sub_font
        for r in (1, 2):
            cell = ws.cell(row=r, column=c)
            cell.fill = fill
            cell.font = font
            cell.alignment = center
            cell.border = border

    # 数据行
    for ri, date in enumerate(all_dates):
        row = 3 + ri
        ws.cell(row=row, column=1, value=date)
        total = 0.0
        for ti, tenor in enumerate(TENOR_ORDER):
            by_inst = data.get((bond, tenor), {})
            c0 = 2 + ti * n_inst
            for ci, inst in enumerate(INST_ORDER):
                v = by_inst.get(inst, {}).get(date)
                cell = ws.cell(row=row, column=c0 + ci)
                if v is not None:
                    cell.value = v
                    total += v
        tc = ws.cell(row=row, column=total_c, value=round(total, 2))
        tc.number_format = "0.00"
        for c in (1, total_c):
            ws.cell(row=row, column=c).border = border
        if total != 0:
            tc.font = Font(color="C00000" if total < 0 else "006100", size=10)

    ws.freeze_panes = "B3"
    ws.column_dimensions["A"].width = 12
    for c in range(2, last_col + 1):
        ws.column_dimensions[get_column_letter(c)].width = 10.5
    ws.row_dimensions[1].height = 18
    ws.row_dimensions[2].height = 18
    return ws


def build_readme(wb, opt, all_dates, missing):
    ws = wb.create_sheet("说明", 0)
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 90
    rows = [
        ("机构行为净买入 · 最细颗粒度全量导出", ""),
        ("", ""),
        ("生成时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
        ("数据来源", "CFETS 现券净买入分机构统计（机构行为监测页同源上游 bondflow API）"),
        ("指标口径", "净买入（亿元）。正值=净买入，负值=净卖出；日度（交易日）数据"),
        ("数据范围", f"{all_dates[0]} ~ {all_dates[-1]}（共 {len(all_dates)} 个交易日）"),
        ("覆盖维度", f"{len(BOND_ORDER)} 券种 × {len(TENOR_ORDER)} 期限档 × {len(INST_ORDER)} 机构类型"),
        ("", ""),
        ("Sheet 结构", "每个券种一个 sheet；行=交易日，列=期限×机构（二列表头），末列为该券种当日全期限全机构合计"),
        ("缺失值", "该机构当日无该券种×期限数据时留空"),
        ("券种清单", "、".join(BOND_ORDER)),
        ("期限档清单", "、".join(TENOR_ORDER)),
        ("机构清单", "、".join(INST_ORDER)),
        ("", ""),
        ("生成脚本", "scripts/export_institution_flow_excel.py（可随时重跑刷新）"),
    ]
    if missing:
        rows.append(("注意", f"以下组合上游无数据（对应单元格留空）：{'、'.join(missing)}"))
    title_font = Font(bold=True, size=14)
    label_font = Font(bold=True, size=10)
    for i, (a, b) in enumerate(rows, start=1):
        ca, cb = ws.cell(row=i, column=1, value=a), ws.cell(row=i, column=2, value=b)
        if i == 1:
            ca.font = title_font
        else:
            ca.font = label_font
        cb.alignment = Alignment(wrap_text=True, vertical="top")


def main():
    print("1/4 拉取维度与最新日期 …")
    opt = fetch_options()
    end_date = opt["latest_date"]
    print(f"    上游数据最新至 {end_date}")

    combos = [(b, t) for b in BOND_ORDER for t in TENOR_ORDER]
    cache = OUT_DIR / "机构行为_拉取缓存.json"
    data = None
    if cache.exists():
        try:
            raw = json.loads(cache.read_text(encoding="utf-8"))
            if raw.get("end_date") == end_date:
                data = {tuple(k.split("|", 1)): v for k, v in raw["data"].items()}
                print(f"    命中本地缓存 {cache}（当日已拉取，直接复用）")
        except Exception:  # noqa: BLE001
            data = None
    if data is None:
        print(f"2/4 并发拉取 {len(combos)} 个（券种×期限）组合 × {len(INST_ORDER)} 机构全量日度序列 …")
        data = {}
        done = 0
        with ThreadPoolExecutor(max_workers=WORKERS) as pool:
            futures = {pool.submit(fetch_combo, b, t, end_date): (b, t) for b, t in combos}
            for fut in as_completed(futures):
                bond, tenor, by_inst = fut.result()
                data[(bond, tenor)] = by_inst
                done += 1
                print(f"    [{done}/{len(combos)}] {bond} × {tenor}")
        try:
            cache.write_text(json.dumps(
                {"end_date": end_date,
                 "data": {f"{b}|{t}": v for (b, t), v in data.items()}},
                ensure_ascii=False), encoding="utf-8")
        except Exception:  # noqa: BLE001
            pass

    all_dates = sorted({d for by_inst in data.values() for s in by_inst.values() for d in s})
    if not all_dates:
        print("上游无任何数据，终止")
        sys.exit(1)
    print(f"    合并后交易日 {len(all_dates)} 天：{all_dates[0]} ~ {all_dates[-1]}")

    missing = [f"{b}×{t}" for (b, t), by_inst in data.items() if not by_inst]
    points = sum(len(s) for by_inst in data.values() for s in by_inst.values())
    print(f"    数据点共 {points} 个，无数据组合 {len(missing)} 个")

    print("3/4 写入 Excel（每个券种一个 sheet）…")
    wb = Workbook()
    wb.remove(wb.active)
    for i, bond in enumerate(BOND_ORDER):
        build_sheet(wb, bond, data, all_dates)
        print(f"    {bond} sheet 完成")
    build_readme(wb, opt, all_dates, missing)

    out = OUT_DIR / f"机构行为净买入_分券种分期限分机构_全量_{datetime.now().strftime('%Y%m%d')}.xlsx"
    print(f"4/4 保存 {out}")
    wb.save(out)
    print(f"完成：{out}（{out.stat().st_size / 1048576:.1f} MB）")


if __name__ == "__main__":
    main()
