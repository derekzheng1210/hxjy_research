# -*- coding: utf-8 -*-
"""
核查：曲线样本中「Excel 票面为空」的个券（此前诊断共 61 只），
从原始 Excel 确认其是否为「取消发行」或「回拨其他期限」。
输出 data/cache/empty_coupon_audit.json + 控制台摘要
"""
import json, os, re, glob, datetime as dt
import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
BASE = os.path.normpath(os.path.join(HERE, ".."))
DATA = os.path.join(BASE, "data")
SRC_DIR = r"D:/2026/一级投标/投标情况"
CURVE = os.path.join(DATA, "coupon_curve.json")
OUT = os.path.join(DATA, "cache", "empty_coupon_audit.json")

# ---------- 1. 读取曲线样本 ----------
curve = json.load(open(CURVE, encoding="utf-8"))
bonds = curve.get("bonds", [])
print(f"曲线样本 {len(bonds)} 只")

# ---------- 2. 读取全部 Excel 行（含全列） ----------
def col_index(header, *keywords):
    for i, name in enumerate(header):
        s = str(name or "")
        if any(k in s for k in keywords):
            return i
    return None

def cell_num(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = re.sub(r"[^\d.\-]", "", str(v).replace(",", ""))
    if not s or s in ("-", "--", "."):
        return None
    try:
        return float(s)
    except ValueError:
        return None

excel_rows = {}   # name -> list of {date, tenor, plan, coupon, forecast, raw...}
files = sorted(glob.glob(os.path.join(SRC_DIR, "一级发行-信用债发行 2026-*.xlsx")))
print(f"Excel 文件 {len(files)} 个")
for path in files:
    date_str = os.path.basename(path)[11:21]
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb.worksheets[0]
    except Exception as e:
        print(f"  跳过 {path}: {e}")
        continue
    header = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    ci_name = col_index(header, "债券简称", "债券名称")
    ci_term = col_index(header, "发行期限", "期限")
    ci_plan = col_index(header, "计划发行")
    ci_coup = col_index(header, "票面利率", "票面")
    ci_fc   = col_index(header, "新债预测", "预测")
    if ci_name is None:
        continue
    for r in range(2, ws.max_row + 1):
        name = str(ws.cell(r, ci_name + 1).value or "").strip()
        if not name:
            continue
        row = {
            "date": date_str,
            "file": os.path.basename(path),
            "row": r,
            "tenor": str(ws.cell(r, ci_term + 1).value or "").strip() if ci_term is not None else None,
            "planYi": cell_num(ws.cell(r, ci_plan + 1).value) if ci_plan is not None else None,
            "coupon": cell_num(ws.cell(r, ci_coup + 1).value) if ci_coup is not None else None,
            "forecast": cell_num(ws.cell(r, ci_fc + 1).value) if ci_fc is not None else None,
            "raw": [ws.cell(r, c).value for c in range(1, ws.max_column + 1)],
            "header": header,
        }
        excel_rows.setdefault(name, []).append(row)

print(f"Excel 主体个券名 {len(excel_rows)} 个")

# ---------- 3. 曲线券 → Excel 匹配，找出票面为空者 ----------
empty_list = []
matched = 0
for b in bonds:
    name = b["name"]
    rows = excel_rows.get(name) or []
    # 优先同簿记日的行
    same = [r for r in rows if r["date"] == b["date"]]
    r = same[0] if same else (rows[0] if rows else None)
    if r is None:
        continue  # Excel 无此券（不属本次核查范围）
    matched += 1
    if r["coupon"] is None or r["coupon"] == 0:
        empty_list.append({"bond": b, "excel": {k: r[k] for k in ("date", "file", "row", "tenor", "planYi", "coupon", "forecast")}})

print(f"曲线券在 Excel 中匹配到 {matched} 只；其中 Excel 票面为空 {len(empty_list)} 只")

# ---------- 4. 同名券全历史追踪（判断回拨：是否在其他日期/其他期限再次出现） ----------
for e in empty_list:
    name = e["bond"]["name"]
    e["allOccurrences"] = [
        {k: r[k] for k in ("date", "tenor", "planYi", "coupon", "forecast")} for r in excel_rows.get(name, [])
    ]

with open(OUT, "w", encoding="utf-8") as f:
    json.dump(empty_list, f, ensure_ascii=False, indent=1)
print("已写入", OUT)

# ---------- 5. 摘要 ----------
print("\n===== 明细（曲线券 | Excel行 | 同名券其他出现）=====")
for e in empty_list:
    b, x = e["bond"], e["excel"]
    occ = e["allOccurrences"]
    others = [o for o in occ if o["date"] != x["date"]]
    print(f"{b['date']} {name if False else b['name']:<24s} DM票面={b['coupon']}% {b['termY']}Y {b['amtYi']}亿")
    print(f"    Excel[{x['date']}] 期限={x['tenor']} 计划={x['planYi']} 票面={x['coupon']} 预测={x['forecast']}")
    for o in others:
        print(f"    其他出现[{o['date']}] 期限={o['tenor']} 计划={o['planYi']} 票面={o['coupon']} 预测={o['forecast']}")
