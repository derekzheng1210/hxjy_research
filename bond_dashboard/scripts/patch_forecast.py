# -*- coding: utf-8 -*-
"""
推荐个券"新债预测"数据修正 v2（修复列错位 bug）：
从每日发行 Excel 的"新债预测"列重新提取（2026-06-15 ~ 08-31）

关键修复：
1. 列定位 bug：部分版式（08-05/08-27/08-31 等）同时有"票面-预测"列(第8列) 和"新债预测"列(第11列)，
   旧逻辑关键词含"预测"会误匹配到"票面-预测"列（存的是 bp 利差如 -12）。
   → 改用精确前缀匹配 ^新债预测，只匹配真正的新债预测列
2. 文本格式数字（'1.67'）也能 float 转换；含空格/全角字符时清洗
"""
import openpyxl, os, json, re, unicodedata

BASE = "D:/2026/一级投标/投标情况"
REC_PATH = "C:/Users/yoyo1/WorkBuddy/2026-09-01-16-51-51/bond-dashboard/data/recommended.json"

rec = json.load(open(REC_PATH, encoding="utf-8"))
by_date = {}
for b in rec:
    by_date.setdefault(b["date"], []).append(b)

def find_forecast_col(ws, max_col=40):
    """精确匹配"新债预测"列（排除"票面-预测"）"""
    for r in range(1, 4):
        for c in range(1, max_col + 1):
            v = ws.cell(r, c).value
            if v is None:
                continue
            s = str(v).strip()
            # 列名以"新债预测"开头（如"新债预测"、"新债预测(%)"）；排除"票面-预测"
            if re.match(r"^新债预测", s):
                return c
    return None

def to_num(v):
    if v is None:
        return None
    # 兼容文本格式数字：去掉不可见字符、单引号、全角空格
    s = str(v).replace("\u00a0", " ").replace("\u3000", " ").strip().strip("'\"")
    if s in ("", "--", "-", "None", "N/A", "#N/A"):
        return None
    try:
        f = float(s.replace(",", "").replace("%", ""))
        return f if f != 0 else None
    except Exception:
        return None

files = [f for f in os.listdir(BASE) if re.match(r"一级发行-信用债发行 2026-\d{2}-\d{2}\.xlsx$", f)]
updated, missing_col = 0, 0
col_map = {}
for fn in files:
    d = fn[11:21]
    recs = by_date.get(d)
    if not recs:
        continue
    wb = openpyxl.load_workbook(os.path.join(BASE, fn), data_only=True)
    ws = wb.worksheets[0]
    cn = None
    for r in range(1, 4):
        for c in range(1, min(ws.max_column, 40) + 1):
            v = ws.cell(r, c).value
            if v and re.match(r"^债券简称", str(v).strip()):
                cn = c
                break
        if cn:
            break
    cf = find_forecast_col(ws)
    if not cn or not cf:
        missing_col += 1
        continue
    m = {b["name"]: b for b in recs}
    for r in range(2, ws.max_row + 1):
        nm = ws.cell(r, cn).value
        if not nm or str(nm).strip() not in m:
            continue
        b = m[str(nm).strip()]
        val = to_num(ws.cell(r, cf).value)
        if val is not None and b.get("forecast") != val:
            b["forecast"] = val
            updated += 1

with open(REC_PATH, "w", encoding="utf-8") as f:
    json.dump(rec, f, ensure_ascii=False, indent=1)

noFc = [b for b in rec if not b["forecast"]]
has = [b for b in rec if b["forecast"]]
print(f"更新预测值: {updated} 只；有预测列文件: {len(files) - missing_col}/{len(files)}；无列文件: {missing_col}")
print(f"有预测: {len(has)} 只；仍无预测: {len(noFc)} 只")
# 校验：之前误存 bp 值的券现在是否修正
for b in rec:
    if b["name"] == "26沪盛K1":
        print(f"校验 26沪盛K1: 新预测={b['forecast']} (之前误存 -12)")
    if b["name"] == "26电网MTN027":
        print(f"校验 26电网MTN027: 新预测={b['forecast']} (之前误存 -5)")
print("有预测样例:")
for b in has[:6]:
    print(f"  {b['date']} {b['name']}: {b['forecast']}")
