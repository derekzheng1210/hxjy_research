# -*- coding: utf-8 -*-
"""
增量更新每日发行 Excel → recommended.json（仅处理日期 > 已导入最大日期的文件）
- 提取「是否推荐=是」的券追加到 recommended.json（date+name 已存在则跳过）
- 精确定位列（^票面 排除 票面-预测 / ^新债预测 / ^票面-预测 / ^YY评分）
- 手工修正过的旧数据完全不动
"""
import openpyxl, os, json, re

BASE = "D:/2026/一级投标/投标情况"
REC_PATH = "C:/Users/yoyo1/WorkBuddy/2026-09-01-16-51-51/bond-dashboard/data/recommended.json"
LAST_DATE = "2026-08-31"  # 已有数据最大日期（含）

rec = json.load(open(REC_PATH, encoding="utf-8"))
existing = {(b["date"], b["name"]) for b in rec}
max_date = max((b["date"] for b in rec), default="2026-06-15")

def find_col(ws, pattern, max_col=45):
    for r in range(1, 4):
        for c in range(1, max_col + 1):
            v = ws.cell(r, c).value
            if v and re.match(pattern, str(v).strip()):
                return c
    return None

def find_coupon_col(ws, max_col=45):
    """票面列：以票面开头但排除票面-预测"""
    for r in range(1, 4):
        for c in range(1, max_col + 1):
            v = ws.cell(r, c).value
            if v:
                s = str(v).strip()
                if re.match(r"^票面", s) and not s.startswith("票面-预测"):
                    return c
    return None

def to_num(v):
    if v is None:
        return None
    s = str(v).replace("\u00a0", " ").replace("\u3000", " ").strip().strip("'\"")
    if s in ("", "--", "-", "None", "N/A", "#N/A"):
        return None
    try:
        f = float(s.replace(",", "").replace("%", ""))
        return f if f != 0 else None
    except Exception:
        return None

files = [f for f in os.listdir(BASE) if re.match(r"一级发行-信用债发行 (2026-\d{2}-\d{2})\.xlsx$", f)]
files = [f for f in files if f[11:21] > max_date]
print(f"待处理新文件: {len(files)} 个: {[f[11:21] for f in files]}")

added, skipped = 0, 0
for fn in sorted(files):
    d = fn[11:21]
    wb = openpyxl.load_workbook(os.path.join(BASE, fn), data_only=True)
    ws = wb.worksheets[0]
    cn = find_col(ws, r"^债券简称")
    ccode = find_col(ws, r"^债券代码")
    cterm = find_col(ws, r"^发行期限")
    cplan = find_col(ws, r"^计划发行")
    ccoupon = find_coupon_col(ws)
    cfc = find_col(ws, r"^新债预测")
    csp = find_col(ws, r"^票面-预测")
    cissuer = find_col(ws, r"^发行人")
    cregion = find_col(ws, r"^区域")
    ctype = find_col(ws, r"^债券类型")
    crec = find_col(ws, r"^是否推荐")
    cyy = find_col(ws, r"^YY评分")
    cpay = find_col(ws, r"^缴款日")
    print(f"\n{d}: 推荐列={crec} 票面列={ccoupon} 新债预测列={cfc} 票面-预测列={csp} YY列={cyy} 缴款列={cpay}")

    for r in range(2, ws.max_row + 1):
        nm = ws.cell(r, cn).value if cn else None
        if not nm:
            continue
        name = str(nm).strip()
        # 是否推荐判断（列存在时只看 = 是）
        if crec:
            rv = ws.cell(r, crec).value
            if str(rv).strip() != "是":
                continue
        if (d, name) in existing:
            skipped += 1
            continue
        cv = ws.cell(r, ccoupon).value if ccoupon else None
        coupon_raw = None
        if cv is None or isinstance(cv, str):
            s = str(cv).strip() if cv is not None else ""
            if s and ("回拨" in s or "取消" in s):
                coupon_raw = s
        coupon = to_num(cv)
        forecast = to_num(ws.cell(r, cfc).value) if cfc else None
        sp_excel = to_num(ws.cell(r, csp).value) if csp else None
        sp_bp = sp_excel if sp_excel is not None else ((round((coupon - forecast) * 100, 1)) if coupon is not None and forecast is not None else None)
        term = str(ws.cell(r, cterm).value).strip() if cterm and ws.cell(r, cterm).value is not None else None
        plan = to_num(ws.cell(r, cplan).value) if cplan else None
        if plan is not None:
            # 计划发行(亿)已是亿元单位
            plan = round(plan, 2)
        yy = None
        if cyy:
            yv = ws.cell(r, cyy).value
            if yv is not None and str(yv).strip():
                yy = str(yv).strip()
        pay_date = None
        if cpay:
            pv = ws.cell(r, cpay).value
            if pv is not None:
                _s = str(pv).strip()
                if _s and _s.lower() != "none":
                    _m = re.search(r"(\d{4})-(\d{2})-(\d{2})", _s)
                    pay_date = _m.group(0) if _m else _s
        item = {
            "date": d,
            "name": name,
            "code": str(ws.cell(r, ccode).value).strip() if ccode and ws.cell(r, ccode).value is not None else None,
            "term": term,
            "plan": plan,
            "actual": plan,
            "coupon": coupon,
            "forecast": forecast,
            "sp_bp": sp_bp,
            "issuer": str(ws.cell(r, cissuer).value).strip() if cissuer and ws.cell(r, cissuer).value is not None else None,
            "region": str(ws.cell(r, cregion).value).strip() if cregion and ws.cell(r, cregion).value is not None else None,
            "bond_type": str(ws.cell(r, ctype).value).strip() if ctype and ws.cell(r, ctype).value is not None else None,
            "yy": yy,
            "pay_date": pay_date,
        }
        if coupon_raw:
            item["coupon_raw"] = coupon_raw
        rec.append(item)
        added += 1
        print(f"  + {d} {name} 票面={coupon} 预测={forecast} sp={sp_bp} YY={yy}")

with open(REC_PATH, "w", encoding="utf-8") as f:
    json.dump(rec, f, ensure_ascii=False, indent=1)
print(f"\n新增 {added} 只，跳过重复 {skipped} 只；recommended.json 现有 {len(rec)} 只")
