# -*- coding: utf-8 -*-
"""回填 YY 评分：每日 Excel"YY评分"列 → recommended.json；登记表"YY评分"列 → bids.json"""
import openpyxl, os, json, re

BASE = "D:/2026/一级投标/投标情况"
DATA = "C:/Users/yoyo1/WorkBuddy/2026-09-01-16-51-51/bond-dashboard/data"

def find_col(ws, pattern, max_col=45):
    for r in range(1, 4):
        for c in range(1, max_col + 1):
            v = ws.cell(r, c).value
            if v and re.match(pattern, str(v).strip()):
                return c
    return None

# ============ 1. recommended.json（每日 Excel YY评分列） ============
rec_path = os.path.join(DATA, "recommended.json")
rec = json.load(open(rec_path, encoding="utf-8"))
by_date = {}
for b in rec:
    by_date.setdefault(b["date"], []).append(b)

updated = 0
files = [f for f in os.listdir(BASE) if re.match(r"一级发行-信用债发行 2026-\d{2}-\d{2}\.xlsx$", f)]
for fn in files:
    d = fn[11:21]
    recs = by_date.get(d)
    if not recs:
        continue
    wb = openpyxl.load_workbook(os.path.join(BASE, fn), data_only=True)
    ws = wb.worksheets[0]
    cn = find_col(ws, r"^债券简称")
    cy = find_col(ws, r"^YY评分|^YY\s*评分")
    if not cn or not cy:
        continue
    m = {b["name"]: b for b in recs}
    for r in range(2, ws.max_row + 1):
        nm = ws.cell(r, cn).value
        if not nm or str(nm).strip() not in m:
            continue
        v = ws.cell(r, cy).value
        if v is not None and str(v).strip() and str(v).strip() != "None":
            yy = str(v).strip()
            if m[str(nm).strip()].get("yy") != yy:
                m[str(nm).strip()]["yy"] = yy
                updated += 1

with open(rec_path, "w", encoding="utf-8") as f:
    json.dump(rec, f, ensure_ascii=False, indent=1)
has_yy = sum(1 for b in rec if b.get("yy"))
print(f"recommended: 更新 {updated}，238 只中有 YY 评分 {has_yy}")

# ============ 2. bids.json（登记表 YY评分列） ============
bids_path = os.path.join(DATA, "store", "bids.json")
bids = json.load(open(bids_path, encoding="utf-8"))
wb = openpyxl.load_workbook(os.path.join(BASE, "一级中标登记表.xlsx"), data_only=True)
ws = wb.worksheets[0]
cn = find_col(ws, r"^债券简称")
cy = find_col(ws, r"^YY评分|^YY\s*评分")
print(f"登记表: 简称列={cn} YY列={cy}")
yy_map = {}
if cn and cy:
    for r in range(2, ws.max_row + 1):
        nm = ws.cell(r, cn).value
        v = ws.cell(r, cy).value
        if nm and v is not None and str(v).strip() and str(nm).strip() not in yy_map:
            yy_map[str(nm).strip()] = str(v).strip()

b_updated = 0
for b in bids:
    yy = yy_map.get(b["bondName"])
    if yy and b.get("yy") != yy:
        b["yy"] = yy
        b_updated += 1
with open(bids_path, "w", encoding="utf-8") as f:
    json.dump(bids, f, ensure_ascii=False, indent=1)
has_yy_b = sum(1 for b in bids if b.get("yy"))
print(f"bids: 更新 {b_updated}，{len(bids)} 条中有 YY 评分 {has_yy_b}")
