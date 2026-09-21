# -*- coding: utf-8 -*-
"""给推荐个券打标记：从每日发行 Excel 票面列原文识别"取消发行"/"回拨XX年期"→ 写入 coupon_raw 字段"""
import openpyxl, os, json, re

BASE = "D:/2026/一级投标/投标情况"
REC_PATH = "C:/Users/yoyo1/WorkBuddy/2026-09-01-16-51-51/bond-dashboard/data/recommended.json"

rec = json.load(open(REC_PATH, encoding="utf-8"))
by_date = {}
for b in rec:
    by_date.setdefault(b["date"], []).append(b)

def find_col(ws, keywords, max_col=40):
    for r in range(1, 4):
        for c in range(1, max_col + 1):
            v = ws.cell(r, c).value
            if v is None:
                continue
            if any(k in str(v) for k in keywords):
                return c
    return None

files = [f for f in os.listdir(BASE) if re.match(r"一级发行-信用债发行 2026-\d{2}-\d{2}\.xlsx$", f)]
updated = 0
for fn in files:
    d = fn[11:21]
    recs = by_date.get(d)
    if not recs:
        continue
    wb = openpyxl.load_workbook(os.path.join(BASE, fn), data_only=True)
    ws = wb.worksheets[0]
    cn = find_col(ws, ["债券简称", "债券名称"])
    cc = find_col(ws, ["票面"])
    if not cn or not cc:
        continue
    # name -> 券
    m = {b["name"]: b for b in recs}
    for r in range(2, ws.max_row + 1):
        nm = ws.cell(r, cn).value
        if not nm or str(nm).strip() not in m:
            continue
        cv = ws.cell(r, cc).value
        b = m[str(nm).strip()]
        if cv is None or isinstance(cv, str):
            raw = str(cv).strip() if cv else None
            # 只标记有意义的状态（回拨/取消），其它非数字原样保留
            if raw and ("回拨" in raw or "取消" in raw or "#VALUE" in raw):
                b["coupon_raw"] = raw
                updated += 1

with open(REC_PATH, "w", encoding="utf-8") as f:
    json.dump(rec, f, ensure_ascii=False, indent=1)

# 统计
flags = [b for b in rec if b.get("coupon_raw")]
print(f"已标记 {len(flags)} 只（更新 {updated}）")
from collections import Counter
kinds = Counter()
for b in flags:
    raw = b["coupon_raw"]
    if "回拨" in raw: kinds["回拨"] += 1
    elif "取消" in raw: kinds["取消发行"] += 1
    else: kinds[raw] += 1
print("分布:", dict(kinds))
for b in flags:
    print(f"  {b['date']} {b['name']}: {b['coupon_raw']}")
